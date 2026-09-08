"""Casos de prueba automatizados del sistema de gestión documental.

Cada prueba corresponde a un caso del documento ``0.4 Pruebas/03_casos_de_prueba.md``
y al requisito que verifica. El identificador CP-xx aparece en el nombre y en la
docstring para mantener la trazabilidad requisito - prueba.
"""

from __future__ import annotations

import io
from pathlib import Path

import pytest

from services import ai_service, procesamiento
from services.extractor import ErrorExtraccion, dividir_en_fragmentos, extraer_texto
from services.rag import similitud_coseno

CORPUS = Path(__file__).resolve().parents[2] / "11_repositorio_pruebas" / "datos_prueba"

CONTRATO = (
    "CONTRATO DE PRESTACION DE SERVICIOS No. CTR-2026-001\n"
    "PARTES: Servicios Andinos Demo S.A.S. y Tecnosoporte Demo Ltda.\n"
    "OBJETO: mantenimiento preventivo de la red de datos.\n"
    "VALOR TOTAL: COP 36.000.000. PLAZO: 12 meses."
)
FACTURA = (
    "FACTURA ELECTRONICA DE VENTA No. FE-2601\n"
    "CLIENTE: Distribuidora Giron Demo S.A.S.\n"
    "SUBTOTAL: COP 4.200.000 IVA (19%): COP 798.000 TOTAL A PAGAR: COP 4.998.000"
)


# ---------------------------------------------------------------------------
# Utilidades
# ---------------------------------------------------------------------------


def crear_repositorio(client, nombre="Archivo corporativo"):
    client.post("/repositorios/nuevo", data={"nombre": nombre, "descripcion": "Demo"},
                follow_redirects=True)
    return 1


def subir(client, repositorio_id, nombre, contenido, titulo="Documento"):
    return client.post(
        f"/repositorios/{repositorio_id}/documentos",
        data={
            "archivo": (io.BytesIO(contenido.encode("utf-8")), nombre),
            "titulo": titulo,
            "descripcion": "",
        },
        content_type="multipart/form-data",
        follow_redirects=True,
    )


# ---------------------------------------------------------------------------
# Autenticación y control de acceso
# ---------------------------------------------------------------------------


def test_cp01_primer_usuario_registrado_es_administrador(client, app):
    """CP-01 (RF-01): el primer registro crea la cuenta de administrador."""
    client.post(
        "/registro",
        data={"nombre": "Administrador", "correo": "admin@demo.com", "clave": "clave123"},
        follow_redirects=True,
    )
    with app.app_context():
        import app as aplicacion

        fila = aplicacion.get_db().execute(
            "SELECT rol FROM usuarios WHERE correo = 'admin@demo.com'"
        ).fetchone()
    assert fila["rol"] == "administrador"


def test_cp02_login_valido_abre_el_panel(client, sesion):
    """CP-02 (RF-01): credenciales correctas dan acceso al panel."""
    respuesta = sesion()
    assert respuesta.status_code == 200
    assert "Indicadores del repositorio" in respuesta.get_data(as_text=True)


def test_cp03_login_invalido_es_rechazado(client, sesion):
    """CP-03 (RNF-04): una contraseña incorrecta no inicia sesión."""
    sesion()
    client.post("/logout", follow_redirects=True)
    respuesta = client.post(
        "/login", data={"correo": "admin@demo.com", "clave": "incorrecta"}, follow_redirects=True
    )
    assert "Correo o contraseña incorrectos." in respuesta.get_data(as_text=True)


def test_cp04_ruta_privada_exige_sesion(client):
    """CP-04 (RNF-05): sin sesión, el panel redirige al formulario de ingreso."""
    respuesta = client.get("/dashboard")
    assert respuesta.status_code == 302
    assert "/login" in respuesta.headers["Location"]


def test_cp05_un_usuario_no_accede_a_documentos_de_otro(client, sesion, app):
    """CP-05 (RF-16): el aislamiento por propietario devuelve 403."""
    sesion(correo="duenio@demo.com")
    crear_repositorio(client)
    subir(client, 1, "contrato.txt", CONTRATO)
    client.post("/logout", follow_redirects=True)

    sesion(correo="otro@demo.com", clave="clave456", nombre="Otro Usuario")
    assert client.get("/documentos/1").status_code == 403


# ---------------------------------------------------------------------------
# Gestión de repositorios y archivos
# ---------------------------------------------------------------------------


def test_cp06_crear_repositorio_y_listarlo(client, sesion):
    """CP-06 (RF-04, RF-05): el repositorio creado aparece en el listado."""
    sesion()
    crear_repositorio(client, "Archivo contable")
    respuesta = client.get("/repositorios")
    assert "Archivo contable" in respuesta.get_data(as_text=True)


def test_cp07_carga_de_txt_queda_procesada(client, sesion, app):
    """CP-07 (RF-06, RF-IA-01): el archivo se guarda y queda en estado Procesado."""
    sesion()
    crear_repositorio(client)
    respuesta = subir(client, 1, "contrato_001.txt", CONTRATO)
    texto = respuesta.get_data(as_text=True)
    assert "Procesado" in texto
    with app.app_context():
        import app as aplicacion

        fila = aplicacion.get_db().execute(
            "SELECT estado_procesamiento, contenido_extraido FROM documentos WHERE id = 1"
        ).fetchone()
    assert fila["estado_procesamiento"] == procesamiento.ESTADO_PROCESADO
    assert "CTR-2026-001" in fila["contenido_extraido"]


def test_cp08_extension_no_permitida_es_rechazada(client, sesion, app):
    """CP-08 (RF-07, RN-07): un .exe se rechaza y queda registrado como ERROR."""
    sesion()
    crear_repositorio(client)
    respuesta = subir(client, 1, "malicioso.exe", "contenido")
    assert "formato no admitido" in respuesta.get_data(as_text=True)
    with app.app_context():
        import app as aplicacion

        db = aplicacion.get_db()
        assert db.execute("SELECT COUNT(*) AS t FROM documentos").fetchone()["t"] == 0
        assert db.execute(
            "SELECT COUNT(*) AS t FROM bitacora WHERE nivel = 'ERROR'"
        ).fetchone()["t"] == 1


def test_cp09_descarga_conserva_el_nombre_original(client, sesion):
    """CP-09 (RF-10, RN-09): la descarga entrega el archivo con su nombre original."""
    sesion()
    crear_repositorio(client)
    subir(client, 1, "contrato_001.txt", CONTRATO)
    respuesta = client.get("/documentos/1/descargar")
    assert respuesta.status_code == 200
    assert "contrato_001.txt" in respuesta.headers["Content-Disposition"]


def test_cp10_eliminar_documento_borra_su_indice(client, sesion, app):
    """CP-10 (RF-11): al eliminar el documento se borran también sus fragmentos."""
    sesion()
    crear_repositorio(client)
    subir(client, 1, "contrato_001.txt", CONTRATO)
    client.post("/documentos/1/eliminar", follow_redirects=True)
    with app.app_context():
        import app as aplicacion

        db = aplicacion.get_db()
        assert db.execute("SELECT COUNT(*) AS t FROM documentos").fetchone()["t"] == 0
        assert db.execute("SELECT COUNT(*) AS t FROM fragmentos").fetchone()["t"] == 0


# ---------------------------------------------------------------------------
# Extracción de contenido
# ---------------------------------------------------------------------------


@pytest.mark.parametrize(
    "ruta_relativa",
    [
        "contratos/contrato_001.txt",
        "facturas/factura_002.docx",
        "informes/informe_003.pdf",
    ],
)
def test_cp11_extraccion_en_los_tres_formatos(ruta_relativa):
    """CP-11 (RF-06): TXT, DOCX y PDF se convierten a texto legible."""
    ruta = CORPUS / ruta_relativa
    if not ruta.exists():
        pytest.skip("Ejecuta scripts/generar_documentos_prueba.py para crear el corpus.")
    texto = extraer_texto(ruta)
    assert len(texto) > 100


def test_cp12_archivo_sin_texto_produce_error_controlado(tmp_path):
    """CP-12 (RNF-07): un archivo vacío lanza ErrorExtraccion, no una excepción cruda."""
    vacio = tmp_path / "vacio.txt"
    vacio.write_text("   ", encoding="utf-8")
    with pytest.raises(ErrorExtraccion):
        extraer_texto(vacio)


# ---------------------------------------------------------------------------
# Procesamiento con IA: clasificación, resumen y extracción
# ---------------------------------------------------------------------------


def test_cp13_clasificacion_resumen_y_extraccion(client, sesion, app):
    """CP-13 (RF-IA-02, RF-IA-03, RF-IA-04): se guardan categoría, resumen y datos."""
    sesion()
    crear_repositorio(client)
    subir(client, 1, "factura_001.txt", FACTURA)
    with app.app_context():
        import app as aplicacion

        fila = aplicacion.get_db().execute(
            "SELECT categoria, resumen, datos_extraidos, modelo_ia FROM documentos WHERE id = 1"
        ).fetchone()
    assert fila["categoria"] == "Factura"
    assert fila["resumen"]
    assert "total" in fila["datos_extraidos"]
    assert fila["modelo_ia"]


def test_cp14_falla_de_ia_deja_estado_error_sin_tumbar_la_app(client, sesion, app, monkeypatch):
    """CP-14 (RF-15, RNF-07): si la API falla, el documento queda en Error con detalle."""

    def falla(*_args, **_kwargs):
        raise ai_service.ErrorIA("HTTP 429: cuota agotada")

    monkeypatch.setattr(ai_service, "analizar_documento", falla)
    sesion()
    crear_repositorio(client)
    respuesta = subir(client, 1, "contrato_001.txt", CONTRATO)
    assert respuesta.status_code == 200
    with app.app_context():
        import app as aplicacion

        fila = aplicacion.get_db().execute(
            "SELECT estado_procesamiento, error_procesamiento, contenido_extraido "
            "FROM documentos WHERE id = 1"
        ).fetchone()
    assert fila["estado_procesamiento"] == procesamiento.ESTADO_ERROR
    assert "cuota agotada" in fila["error_procesamiento"]
    # El texto extraído se conserva aunque la IA falle.
    assert "CTR-2026-001" in fila["contenido_extraido"]


def test_cp15_reprocesar_recupera_un_documento_con_error(client, sesion, app, monkeypatch):
    """CP-15 (RF-IA-01): reprocesar un documento en Error lo deja Procesado."""
    from tests.conftest import analisis_simulado

    monkeypatch.setattr(
        ai_service, "analizar_documento", lambda *a, **k: (_ for _ in ()).throw(
            ai_service.ErrorIA("servicio no disponible")
        )
    )
    sesion()
    crear_repositorio(client)
    subir(client, 1, "contrato_001.txt", CONTRATO)

    monkeypatch.setattr(ai_service, "analizar_documento", analisis_simulado)
    respuesta = client.post("/documentos/1/procesar", follow_redirects=True)
    assert respuesta.status_code == 200
    with app.app_context():
        import app as aplicacion

        fila = aplicacion.get_db().execute(
            "SELECT estado_procesamiento FROM documentos WHERE id = 1"
        ).fetchone()
    assert fila["estado_procesamiento"] == procesamiento.ESTADO_PROCESADO


# ---------------------------------------------------------------------------
# Búsqueda y consulta en lenguaje natural
# ---------------------------------------------------------------------------


def test_cp16_busqueda_encuentra_texto_dentro_del_archivo(client, sesion):
    """CP-16 (RF-12): la búsqueda recorre el contenido, no solo los metadatos."""
    sesion()
    crear_repositorio(client)
    subir(client, 1, "contrato_001.txt", CONTRATO, titulo="Documento uno")
    respuesta = client.get("/buscar?q=Tecnosoporte&modo=texto")
    texto = respuesta.get_data(as_text=True)
    assert "1 resultado" in texto
    assert "Documento uno" in texto


def test_cp17_busqueda_sin_coincidencias_no_falla(client, sesion):
    """CP-17 (caso límite): un término inexistente devuelve cero resultados."""
    sesion()
    crear_repositorio(client)
    subir(client, 1, "contrato_001.txt", CONTRATO)
    respuesta = client.get("/buscar?q=zzzzz&modo=texto")
    assert "No se encontraron coincidencias" in respuesta.get_data(as_text=True)


def test_cp18_busqueda_semantica_ordena_por_similitud(client, sesion):
    """CP-18 (RF-IA-05): la búsqueda semántica devuelve primero el documento afín."""
    sesion()
    crear_repositorio(client)
    subir(client, 1, "contrato_001.txt", CONTRATO, titulo="Contrato de red")
    subir(client, 1, "factura_001.txt", FACTURA, titulo="Factura de licencia")
    respuesta = client.get("/buscar?q=cuanto%20es%20el%20iva%20de%20la%20factura&modo=semantica")
    texto = respuesta.get_data(as_text=True)
    assert texto.index("Factura de licencia") < texto.index("Contrato de red")


def test_cp19_asistente_responde_y_registra_la_consulta(client, sesion, app):
    """CP-19 (RF-IA-05): la pregunta en lenguaje natural devuelve respuesta y fuentes."""
    sesion()
    crear_repositorio(client)
    subir(client, 1, "factura_001.txt", FACTURA, titulo="Factura de licencia")
    respuesta = client.post(
        "/asistente", data={"pregunta": "¿Cuál es el total de la factura?"}, follow_redirects=True
    )
    texto = respuesta.get_data(as_text=True)
    assert "Respuesta simulada" in texto
    assert "Fragmentos utilizados" in texto
    with app.app_context():
        import app as aplicacion

        fila = aplicacion.get_db().execute(
            "SELECT pregunta, fuentes FROM consultas_ia"
        ).fetchone()
    assert "total" in fila["pregunta"].lower()
    assert "Factura de licencia" in fila["fuentes"]


def test_cp20_asistente_sin_documentos_avisa(client, sesion):
    """CP-20 (caso límite): preguntar sin documentos indexados muestra un aviso."""
    sesion()
    respuesta = client.post(
        "/asistente", data={"pregunta": "¿Qué contratos hay?"}, follow_redirects=True
    )
    assert "no hay documentos indexados" in respuesta.get_data(as_text=True).lower()


# ---------------------------------------------------------------------------
# Panel, bitácora y utilidades
# ---------------------------------------------------------------------------


def test_cp21_dashboard_muestra_indicadores(client, sesion):
    """CP-21 (RF-14): el panel presenta totales, categorías y estados."""
    sesion()
    crear_repositorio(client)
    subir(client, 1, "factura_001.txt", FACTURA)
    texto = client.get("/dashboard").get_data(as_text=True)
    assert "Procesados por IA" in texto
    assert "Documentos por categoría" in texto
    assert "Factura" in texto


def test_cp22_bitacora_registra_las_operaciones(client, sesion, app):
    """CP-22 (RF-15, RNF-12): cada operación relevante queda en la bitácora."""
    sesion()
    crear_repositorio(client)
    subir(client, 1, "contrato_001.txt", CONTRATO)
    with app.app_context():
        import app as aplicacion

        acciones = {
            fila["accion"]
            for fila in aplicacion.get_db().execute("SELECT accion FROM bitacora")
        }
    assert {"CREAR_REPOSITORIO", "CARGAR_DOCUMENTO", "PROCESAR_DOCUMENTO"} <= acciones


def test_cp23_recurso_inexistente_devuelve_404(client, sesion):
    """CP-23 (caso límite): un identificador que no existe muestra la página 404."""
    sesion()
    assert client.get("/documentos/999").status_code == 404


def test_cp24_chequeo_de_salud_responde(client):
    """CP-24 (RF-17): el endpoint de salud informa el estado del servicio y de la IA."""
    datos = client.get("/salud").get_json()
    assert datos["estado"] == "ok"
    assert "modelo_texto" in datos["ia"]


# ---------------------------------------------------------------------------
# Pruebas unitarias de los algoritmos de apoyo
# ---------------------------------------------------------------------------


def test_cp25_fragmentacion_respeta_tamano_y_solape():
    """CP-25: el texto largo se divide en fragmentos con solape."""
    texto = ("Frase de prueba numero uno. " * 200).strip()
    fragmentos = dividir_en_fragmentos(texto, tamano=500, solape=50)
    assert len(fragmentos) > 1
    assert all(len(f) <= 500 for f in fragmentos)


def test_cp26_similitud_coseno_es_coherente():
    """CP-26: vectores iguales dan 1.0 y vectores ortogonales dan 0.0."""
    assert similitud_coseno([1, 0, 1], [1, 0, 1]) == pytest.approx(1.0)
    assert similitud_coseno([1, 0], [0, 1]) == pytest.approx(0.0)
    assert similitud_coseno([], [1, 2]) == 0.0


# ---------------------------------------------------------------------------
# Carga múltiple, formatos adicionales y cola de procesamiento
# ---------------------------------------------------------------------------


def subir_varios(client, ruta, archivos, datos=None):
    """Envía varios archivos en un mismo formulario."""
    campos = {"archivo": [(io.BytesIO(c.encode("utf-8")), n) for n, c in archivos]}
    campos.update(datos or {})
    return client.post(
        ruta, data=campos, content_type="multipart/form-data", follow_redirects=True
    )


def test_cp30_carga_multiple_desde_el_asistente(client, sesion, app):
    """CP-30 (RF-06): el asistente acepta varios archivos y crea el repositorio."""
    sesion()
    respuesta = subir_varios(
        client,
        "/asistente/documentos",
        [("contrato.txt", CONTRATO), ("factura.txt", FACTURA), ("notas.md", "# " + CONTRATO)],
    )
    assert "3 documento(s) cargados" in respuesta.get_data(as_text=True)
    with app.app_context():
        import app as aplicacion

        db = aplicacion.get_db()
        assert db.execute("SELECT COUNT(*) AS t FROM documentos").fetchone()["t"] == 3
        assert db.execute(
            "SELECT COUNT(*) AS t FROM documentos WHERE estado_procesamiento = 'Procesado'"
        ).fetchone()["t"] == 3
        # El repositorio se creó solo, sin pedirlo al usuario.
        assert db.execute(
            "SELECT nombre FROM repositorios"
        ).fetchone()["nombre"] == "Documentos del asistente"


def test_cp31_un_archivo_invalido_no_cancela_el_lote(client, sesion, app):
    """CP-31 (RF-07): los archivos válidos se cargan aunque otro sea rechazado."""
    sesion()
    respuesta = subir_varios(
        client,
        "/asistente/documentos",
        [("bueno.txt", CONTRATO), ("malo.exe", "binario"), ("otro.txt", FACTURA)],
    )
    texto = respuesta.get_data(as_text=True)
    assert "2 documento(s) cargados" in texto
    assert "no se cargaron" in texto
    with app.app_context():
        import app as aplicacion

        assert aplicacion.get_db().execute(
            "SELECT COUNT(*) AS t FROM documentos"
        ).fetchone()["t"] == 2


@pytest.mark.parametrize(
    "nombre, contenido, esperado",
    [
        ("reporte.html", "<html><body><h1>Informe anual</h1><p>Ventas del periodo</p></body></html>", "Informe anual"),
        ("datos.csv", "cliente,total\nDistribuidora Demo,4998000\nCooperativa Demo,120000", "Distribuidora Demo"),
        ("notas.md", "# Contrato marco\n\nValor total: COP 36.000.000 a doce meses.", "Contrato marco"),
        ("config.json", '{"proveedor": "Tecnosoporte Demo", "valor": 36000000}', "Tecnosoporte Demo"),
    ],
)
def test_cp32_formatos_de_texto_adicionales(tmp_path, nombre, contenido, esperado):
    """CP-32 (RF-06): HTML, CSV, Markdown y JSON se convierten a texto."""
    ruta = tmp_path / nombre
    ruta.write_text(contenido, encoding="utf-8")
    texto = extraer_texto(ruta)
    assert esperado in texto


def test_cp33_documento_word_moderno_renombrado_como_doc(tmp_path):
    """CP-33 (RF-06): un .docx con extensión .doc se detecta por su firma."""
    from docx import Document

    documento = Document()
    documento.add_paragraph("CONTRATO DE PRESTACION DE SERVICIOS No. CTR-2026-050")
    documento.add_paragraph("VALOR TOTAL: COP 12.500.000 a seis meses.")
    ruta = tmp_path / "contrato_antiguo.doc"
    documento.save(str(ruta))

    texto = extraer_texto(ruta)
    assert "CTR-2026-050" in texto


def test_cp34_ocr_de_respaldo_para_archivos_sin_texto(client, sesion, app, monkeypatch):
    """CP-34 (RF-IA-01): si no hay texto local, la IA transcribe el archivo."""
    llamadas = []

    def transcribir(ruta):
        llamadas.append(str(ruta))
        return FACTURA

    monkeypatch.setattr(ai_service, "transcribir_archivo", transcribir)
    sesion()
    crear_repositorio(client)
    # Un PNG no tiene extracción local: debe caer en la transcripción con IA.
    client.post(
        "/repositorios/1/documentos",
        data={"archivo": (io.BytesIO(b"\x89PNG\r\n\x1a\n" + b"0" * 200), "factura.png")},
        content_type="multipart/form-data",
        follow_redirects=True,
    )
    assert len(llamadas) == 1
    with app.app_context():
        import app as aplicacion

        fila = aplicacion.get_db().execute(
            "SELECT estado_procesamiento, categoria FROM documentos WHERE id = 1"
        ).fetchone()
    assert fila["estado_procesamiento"] == procesamiento.ESTADO_PROCESADO
    assert fila["categoria"] == "Factura"


def test_cp35_la_cola_procesa_los_documentos_pendientes(client, sesion, app):
    """CP-35 (RF-IA-01): el trabajador drena los documentos en estado Pendiente."""
    sesion()
    crear_repositorio(client)
    with app.app_context():
        import app as aplicacion

        db = aplicacion.get_db()
        db.execute(
            """
            INSERT INTO documentos (
                repositorio_id, titulo, descripcion, categoria, nombre_original,
                nombre_almacenado, tipo_mime, tamano_bytes, estado_procesamiento, creado_por
            ) VALUES (1, 'Pendiente', '', 'Otro', 'pendiente.txt', 'pendiente.txt',
                      'text/plain', 10, 'Pendiente', 1)
            """
        )
        db.commit()
        ruta = Path(app.config["UPLOAD_FOLDER"]) / "pendiente.txt"
        ruta.write_text(CONTRATO, encoding="utf-8")
        assert procesamiento.contar_pendientes(db) == 1

    resumen = procesamiento.procesar_pendientes(
        app.config["DATABASE"], app.config["UPLOAD_FOLDER"]
    )
    assert resumen == {"procesados": 1, "fallidos": 0}

    with app.app_context():
        import app as aplicacion

        db = aplicacion.get_db()
        assert procesamiento.contar_pendientes(db) == 0
        assert db.execute(
            "SELECT categoria FROM documentos WHERE nombre_original = 'pendiente.txt'"
        ).fetchone()["categoria"] == "Contrato"


def test_cp36_archivo_que_supera_el_limite_por_archivo(client, sesion, app):
    """CP-36 (RF-07, RN-08): se rechaza el archivo que excede el tamaño máximo."""
    app.config["MAX_FILE_SIZE"] = 100
    sesion()
    crear_repositorio(client)
    respuesta = subir(client, 1, "grande.txt", "x" * 500)
    assert "supera los" in respuesta.get_data(as_text=True)
    with app.app_context():
        import app as aplicacion

        assert aplicacion.get_db().execute(
            "SELECT COUNT(*) AS t FROM documentos"
        ).fetchone()["t"] == 0
