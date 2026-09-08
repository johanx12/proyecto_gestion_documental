"""Genera el repositorio de 30 documentos de prueba en TXT, DOCX y PDF.

Los datos son ficticios: empresas, valores y fechas fueron inventados para la
práctica académica y no corresponden a personas u organizaciones reales.

Uso:
    python scripts/generar_documentos_prueba.py

Requiere ``reportlab`` para la salida en PDF (ver requirements-dev.txt).
"""

from __future__ import annotations

import sys
from pathlib import Path

from docx import Document
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

DESTINO = Path(__file__).resolve().parents[2] / "11_repositorio_pruebas" / "datos_prueba"

AVISO = (
    "Documento ficticio elaborado para la práctica académica de gestión "
    "documental. No contiene datos personales ni información real."
)

PROVEEDORES = [
    ("Tecnosoporte Demo Ltda.", "901.555.888-1", "mantenimiento preventivo y correctivo de la red de datos"),
    ("Papelería Andina Demo S.A.S.", "900.222.333-4", "suministro de papelería y útiles de oficina"),
    ("Vigilancia Cordillera Demo S.A.", "900.444.121-9", "servicio de vigilancia física en la sede principal"),
    ("Aseo Integral Demo S.A.S.", "901.777.010-2", "aseo y cafetería para las tres sedes administrativas"),
    ("Transportes del Oriente Demo", "900.909.808-7", "transporte terrestre de mercancía entre bodegas"),
    ("Consultores Fiscales Demo S.A.S.", "901.313.414-5", "asesoría contable y tributaria mensual"),
    ("Software Bucaramanga Demo", "901.616.717-8", "desarrollo y soporte del portal de clientes"),
    ("Capacitación Empresarial Demo", "900.818.919-0", "formación en seguridad y salud en el trabajo"),
    ("Mantenimiento Eléctrico Demo", "901.020.121-6", "mantenimiento de subestaciones y plantas eléctricas"),
    ("Publicidad Santander Demo Ltda.", "900.303.404-3", "diseño y producción de material publicitario"),
]

VALORES = [36000000, 12500000, 84000000, 27600000, 45200000,
           18900000, 96400000, 15300000, 52800000, 23700000]
PLAZOS = [12, 6, 24, 12, 18, 12, 24, 4, 12, 8]

CLIENTES = [
    ("Distribuidora Girón Demo S.A.S.", "900.121.232-1"),
    ("Almacenes Floridablanca Demo", "901.343.454-2"),
    ("Cooperativa Piedecuesta Demo", "900.565.676-3"),
    ("Industrias Barrancabermeja Demo", "901.787.898-4"),
    ("Comercial Socorro Demo S.A.S.", "900.909.010-5"),
    ("Agro San Gil Demo Ltda.", "901.121.232-6"),
    ("Hotelera Málaga Demo", "900.343.454-7"),
    ("Constructora Lebrija Demo", "901.565.676-8"),
    ("Textiles Rionegro Demo", "900.787.898-9"),
    ("Logística Zapatoca Demo", "901.909.010-0"),
]

ITEMS = [
    ("Licencia anual de software de gestión", 1, 4200000),
    ("Resma de papel carta (caja x 10)", 25, 138000),
    ("Servicio de mantenimiento mensual", 3, 1850000),
    ("Computador portátil corporativo", 4, 3450000),
    ("Silla ergonómica de oficina", 12, 620000),
    ("Hora de consultoría especializada", 40, 185000),
    ("Impresora multifuncional láser", 2, 2790000),
    ("Kit de aseo institucional", 30, 96000),
    ("Cámara de seguridad IP", 8, 540000),
    ("Póliza de soporte extendido", 1, 7300000),
]

AREAS = [
    ("Gestión documental", "digitalización del archivo central"),
    ("Talento humano", "rotación de personal y clima laboral"),
    ("Financiera", "ejecución presupuestal y cartera"),
    ("Operaciones", "cumplimiento de entregas y tiempos de despacho"),
    ("Tecnología", "disponibilidad de los servicios y soporte"),
    ("Comercial", "ventas por línea de producto"),
    ("Calidad", "hallazgos de auditoría interna"),
    ("Logística", "rotación de inventario en bodega"),
    ("Servicio al cliente", "peticiones, quejas y reclamos"),
    ("Seguridad y salud", "accidentalidad y ausentismo"),
]

TRIMESTRES = ["I trimestre de 2026", "II trimestre de 2026", "III trimestre de 2026",
              "IV trimestre de 2025", "I trimestre de 2026", "II trimestre de 2026",
              "III trimestre de 2026", "IV trimestre de 2025", "I trimestre de 2026",
              "II trimestre de 2026"]


def pesos(valor: int) -> str:
    return f"COP {valor:,.0f}".replace(",", ".")


def contrato(i: int) -> tuple[str, list[str]]:
    proveedor, nit, objeto = PROVEEDORES[i]
    valor = VALORES[i]
    plazo = PLAZOS[i]
    numero = f"CTR-2026-{i + 1:03d}"
    lineas = [
        f"CONTRATO DE PRESTACION DE SERVICIOS No. {numero}",
        "",
        "PARTES: Servicios Andinos Demo S.A.S., NIT 900.111.222-3, en calidad de "
        f"contratante, y {proveedor}, NIT {nit}, en calidad de contratista.",
        f"OBJETO: {objeto.capitalize()}.",
        f"VALOR TOTAL: {pesos(valor)} incluido IVA, pagaderos en cuotas mensuales iguales.",
        f"PLAZO DE EJECUCION: {plazo} meses contados a partir del acta de inicio.",
        f"FECHA DE FIRMA: {(i % 28) + 1:02d} de febrero de 2026, Bucaramanga.",
        "SUPERVISION: La ejecución será supervisada por la Dirección Administrativa.",
        "CLAUSULA PENAL: Equivalente al diez por ciento (10%) del valor del contrato.",
        "TERMINACION: Por vencimiento del plazo, mutuo acuerdo o incumplimiento grave.",
        "",
        AVISO,
    ]
    return f"contrato_{i + 1:03d}", lineas


def factura(i: int) -> tuple[str, list[str]]:
    cliente, nit_cliente = CLIENTES[i]
    descripcion, cantidad, unitario = ITEMS[i]
    subtotal = cantidad * unitario
    iva = round(subtotal * 0.19)
    total = subtotal + iva
    numero = f"FE-{2600 + i + 1}"
    lineas = [
        f"FACTURA ELECTRONICA DE VENTA No. {numero}",
        "",
        "EMISOR: Servicios Andinos Demo S.A.S. - NIT 900.111.222-3 - Bucaramanga, Santander.",
        f"CLIENTE: {cliente} - NIT {nit_cliente}.",
        f"FECHA DE EMISION: {(i % 28) + 1:02d} de marzo de 2026.",
        f"FORMA DE PAGO: Credito a {30 if i % 2 == 0 else 45} dias.",
        "",
        "DETALLE:",
        f"  {cantidad} x {descripcion} - valor unitario {pesos(unitario)}",
        "",
        f"SUBTOTAL: {pesos(subtotal)}",
        f"IVA (19%): {pesos(iva)}",
        f"TOTAL A PAGAR: {pesos(total)}",
        "",
        AVISO,
    ]
    return f"factura_{i + 1:03d}", lineas


def informe(i: int) -> tuple[str, list[str]]:
    area, tema = AREAS[i]
    periodo = TRIMESTRES[i]
    avance = 60 + (i * 3) % 35
    lineas = [
        f"INFORME DE GESTION - AREA DE {area.upper()}",
        "",
        f"PERIODO EVALUADO: {periodo}.",
        f"RESPONSABLE: Jefatura de {area}.",
        f"OBJETIVO: Presentar los resultados sobre {tema}.",
        "",
        "HALLAZGOS:",
        f"  1. El indicador principal del area alcanzo un cumplimiento del {avance}%.",
        f"  2. Se identificaron {2 + i % 4} oportunidades de mejora en los procesos revisados.",
        f"  3. El tiempo promedio de respuesta bajo a {3 + i % 5} dias habiles.",
        "",
        "CONCLUSION: El area mantiene la tendencia esperada, aunque se recomienda "
        "reforzar el seguimiento mensual y documentar los procedimientos criticos.",
        "RECOMENDACION: Asignar un responsable de seguimiento y revisar el plan en el "
        "siguiente comite directivo.",
        "",
        AVISO,
    ]
    return f"informe_{i + 1:03d}", lineas


def guardar_txt(ruta: Path, lineas: list[str]) -> None:
    ruta.write_text("\n".join(lineas), encoding="utf-8")


def guardar_docx(ruta: Path, lineas: list[str]) -> None:
    documento = Document()
    documento.add_heading(lineas[0], level=1)
    for linea in lineas[1:]:
        documento.add_paragraph(linea)
    documento.save(str(ruta))


def guardar_pdf(ruta: Path, lineas: list[str]) -> None:
    estilos = getSampleStyleSheet()
    pdf = SimpleDocTemplate(str(ruta), pagesize=letter, title=lineas[0])
    bloques = [Paragraph(lineas[0], estilos["Heading2"]), Spacer(1, 10)]
    for linea in lineas[1:]:
        bloques.append(Paragraph(linea if linea else "&nbsp;", estilos["BodyText"]))
    pdf.build(bloques)


def main() -> int:
    generadores = {
        "contratos": contrato,
        "facturas": factura,
        "informes": informe,
    }
    # Los formatos se alternan para que el repositorio tenga PDF, DOCX y TXT.
    escritores = [(".txt", guardar_txt), (".docx", guardar_docx), (".pdf", guardar_pdf)]

    total = 0
    for carpeta, generador in generadores.items():
        destino = DESTINO / carpeta
        destino.mkdir(parents=True, exist_ok=True)
        for anterior in destino.iterdir():
            if anterior.is_file():
                anterior.unlink()
        for i in range(10):
            nombre, lineas = generador(i)
            extension, escritor = escritores[i % 3]
            escritor(destino / f"{nombre}{extension}", lineas)
            total += 1
        print(f"{carpeta}: 10 documentos")

    print(f"Total generado: {total} documentos en {DESTINO}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
